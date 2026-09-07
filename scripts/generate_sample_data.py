"""Generate deterministic sample data for the CommercialOps Control Tower."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


OUTPUT_PATH = Path(__file__).resolve().parents[1] / "data" / "commercial_ops_data.xlsx"


def create_sample_data() -> dict[str, pd.DataFrame]:
    """Return the five sample commercial operations datasets."""
    customers = pd.DataFrame(
        [
            ["CUST-001", "Northstar Retail", "North", "Enterprise", "ops@northstar.example"],
            ["CUST-002", "BluePeak Foods", "West", "Mid-Market", "finance@bluepeak.example"],
            ["CUST-003", "Greenline Logistics", "South", "Enterprise", "billing@greenline.example"],
            ["CUST-004", "Aster Health", "East", "Mid-Market", "accounts@aster.example"],
            ["CUST-005", "Summit Textiles", "West", "SMB", "owner@summit.example"],
            ["CUST-006", "Cedar Manufacturing", "North", "Enterprise", "ap@cedar.example"],
            ["CUST-007", "Harbor Hospitality", "South", "Mid-Market", "finance@harbor.example"],
            ["CUST-008", "Orbit Learning", "East", "SMB", "admin@orbitlearning.example"],
            ["CUST-009", "MetroBuild Services", "West", "Enterprise", "payments@metrobuild.example"],
            ["CUST-010", "Sunrise Energy", "South", "Enterprise", "procurement@sunrise.example"],
            ["CUST-011", "Willow Consulting", "North", "SMB", "hello@willow.example"],
            ["CUST-012", "Cobalt Media", "East", "Mid-Market", "billing@cobalt.example"],
            ["CUST-013", "Redwood Pharma", "West", "Enterprise", "ap@redwood.example"],
            ["CUST-014", "Nexa Telecom", "South", "Mid-Market", None],
            ["CUST-004", "Aster Health", "East", "Mid-Market", "accounts@aster.example"],
        ],
        columns=["customer_id", "customer_name", "region", "segment", "email"],
    )

    orders = pd.DataFrame(
        [
            ["ORD-001", "CUST-001", "2026-01-08", 25000, "PRJ-001", "Completed"],
            ["ORD-002", "CUST-002", "2026-01-19", 42000, "PRJ-002", "Completed"],
            ["ORD-003", "CUST-003", "2026-02-03", 18000, "PRJ-003", "Completed"],
            ["ORD-004", "CUST-004", "2026-02-17", 65000, "PRJ-004", "In Progress"],
            ["ORD-005", "CUST-005", "2026-03-02", 12000, "PRJ-006", "Completed"],
            ["ORD-006", "CUST-006", "2026-03-14", 33000, "PRJ-001", "Completed"],
            ["ORD-007", "CUST-007", "2026-03-28", 28500, "PRJ-007", "Completed"],
            ["ORD-008", "CUST-008", "2026-04-05", 76000, "PRJ-004", "In Progress"],
            ["ORD-009", "CUST-009", "2026-04-18", 54000, "PRJ-005", "Completed"],
            ["ORD-010", "CUST-010", "2026-04-29", 22000, "PRJ-003", "Completed"],
            ["ORD-011", "CUST-011", "2026-05-07", 48000, "PRJ-002", "Completed"],
            ["ORD-012", "CUST-012", "2026-05-19", 19500, "PRJ-006", "Completed"],
            ["ORD-013", "CUST-013", "2026-05-31", 88000, "PRJ-005", "Completed"],
            ["ORD-014", "CUST-014", "2026-06-09", 31000, "PRJ-008", "Completed"],
            ["ORD-015", "CUST-001", "2026-06-16", 14500, "PRJ-007", "Completed"],
            ["ORD-016", "CUST-002", "2026-06-27", 67000, "PRJ-004", "In Progress"],
            ["ORD-017", "CUST-003", "2026-07-04", 26000, "PRJ-001", "Completed"],
            ["ORD-018", "CUST-004", "2026-07-12", 39000, "PRJ-002", "Completed"],
            ["ORD-019", "CUST-006", "2026-07-21", 52000, "PRJ-005", "Completed"],
            ["ORD-020", "CUST-007", "2026-07-30", 23500, "PRJ-006", "In Progress"],
            ["ORD-021", "CUST-008", "2026-08-05", 46000, "PRJ-003", "In Progress"],
            ["ORD-022", "CUST-009", "2026-08-11", 71000, "PRJ-005", "In Progress"],
            ["ORD-023", "CUST-011", "2026-08-18", 16500, "PRJ-007", "On Hold"],
            ["ORD-024", "CUST-013", "2026-08-24", 93000, "PRJ-008", "In Progress"],
            ["ORD-025", "CUST-999", "2026-08-29", 27500, "PRJ-001", "In Progress"],
        ],
        columns=["order_id", "customer_id", "order_date", "order_amount", "project_id", "status"],
    )
    orders["order_date"] = pd.to_datetime(orders["order_date"])

    invoices = pd.DataFrame(
        [
            ["INV-001", "ORD-001", "2026-01-10", "2026-02-09", 25000, "Paid"],
            ["INV-002", "ORD-002", "2026-01-22", "2026-02-21", 42000, "Paid"],
            ["INV-003", "ORD-003", "2026-02-05", "2026-03-07", 18000, "Paid"],
            ["INV-004", "ORD-004", "2026-02-20", "2026-03-22", 65000, "Partially Paid"],
            ["INV-005", "ORD-005", "2026-03-04", "2026-04-03", 15000, "Paid"],
            ["INV-006", "ORD-006", "2026-03-17", "2026-04-16", 33000, "Overdue"],
            ["INV-007", "ORD-007", "2026-03-31", "2026-04-30", 28500, "Paid"],
            ["INV-008", "ORD-008", "2026-04-08", "2026-05-08", 76000, "Partially Paid"],
            ["INV-009", "ORD-009", "2026-04-21", "2026-05-21", 54000, "Paid"],
            ["INV-010", "ORD-010", "2026-05-02", "2026-06-01", 22000, "Overdue"],
            ["INV-011", "ORD-011", "2026-05-10", "2026-06-09", 48000, "Paid"],
            ["INV-012", "ORD-012", "2026-05-22", "2026-06-21", 19500, "Overdue"],
            ["INV-013", "ORD-013", "2026-06-03", "2026-07-03", 88000, "Paid"],
            ["INV-014", "ORD-014", "2026-06-12", "2026-07-12", 31000, "Paid"],
            ["INV-015", "ORD-015", "2026-06-19", "2026-07-19", 14500, "Partially Paid"],
            ["INV-016", "ORD-016", "2026-06-30", "2026-07-30", 67000, "Overdue"],
            ["INV-017", "ORD-017", "2026-07-07", "2026-08-06", 26000, "Paid"],
            ["INV-018", "ORD-018", "2026-07-15", "2026-08-14", 39000, "Paid"],
            ["INV-019", "ORD-019", "2026-08-02", "2026-09-30", 52000, "Unpaid"],
            ["INV-020", "ORD-999", "2026-08-20", "2026-09-19", 17500, "Unpaid"],
        ],
        columns=["invoice_id", "order_id", "invoice_date", "due_date", "invoice_amount", "status"],
    )
    invoices[["invoice_date", "due_date"]] = invoices[["invoice_date", "due_date"]].apply(
        pd.to_datetime
    )

    payments = pd.DataFrame(
        [
            ["PAY-001", "INV-001", "2026-02-02", 25000, "Bank Transfer"],
            ["PAY-002", "INV-002", "2026-02-18", 42000, "Bank Transfer"],
            ["PAY-003", "INV-003", "2026-03-05", 18000, "Credit Card"],
            ["PAY-004", "INV-004", "2026-03-18", 30000, "Bank Transfer"],
            ["PAY-005", "INV-005", "2026-03-29", 15000, "Credit Card"],
            ["PAY-006", "INV-007", "2026-04-27", 28500, "Bank Transfer"],
            ["PAY-007", "INV-008", "2026-05-04", 40000, "Bank Transfer"],
            ["PAY-008", "INV-008", "2026-06-02", 10000, "Bank Transfer"],
            ["PAY-009", "INV-009", "2026-05-17", 54000, "NEFT"],
            ["PAY-010", "INV-011", "2026-06-04", 48000, "NEFT"],
            ["PAY-011", "INV-013", "2026-06-29", 88000, "Bank Transfer"],
            ["PAY-012", "INV-014", "2026-07-08", 31000, "Credit Card"],
            ["PAY-013", "INV-015", "2026-07-15", 5000, "Credit Card"],
            ["PAY-014", "INV-017", "2026-08-02", 26000, "NEFT"],
            ["PAY-015", "INV-018", "2026-08-10", 39000, "Bank Transfer"],
        ],
        columns=["payment_id", "invoice_id", "payment_date", "payment_amount", "payment_method"],
    )
    payments["payment_date"] = pd.to_datetime(payments["payment_date"])

    projects = pd.DataFrame(
        [
            ["PRJ-001", "North Region Rollout", "Aarav Mehta", 180000, 164000, "Active"],
            ["PRJ-002", "Finance Process Upgrade", "Maya Rao", 145000, 138500, "Completed"],
            ["PRJ-003", "Customer Portal Launch", "Vikram Shah", 120000, 113000, "Active"],
            ["PRJ-004", "Enterprise Integration", "Nisha Kapoor", 240000, 226000, "Active"],
            ["PRJ-005", "Warehouse Transformation", "Rohan Iyer", 150000, 247500, "At Risk"],
            ["PRJ-006", "Billing Automation", "Leena Das", 95000, 89000, "Completed"],
            ["PRJ-007", "Service Desk Refresh", "Kabir Singh", 80000, 73500, "Active"],
            ["PRJ-008", "Partner Enablement", None, 130000, 61000, "Active"],
        ],
        columns=["project_id", "project_name", "owner", "budget", "actual_cost", "status"],
    )

    return {
        "Customers": customers,
        "Orders": orders,
        "Invoices": invoices,
        "Payments": payments,
        "Projects": projects,
    }


def format_workbook(path: Path) -> None:
    """Apply simple, consistent formatting to the generated workbook."""
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    currency_columns = {
        "Orders": {"order_amount"},
        "Invoices": {"invoice_amount"},
        "Payments": {"payment_amount"},
        "Projects": {"budget", "actual_cost"},
    }
    date_columns = {
        "Orders": {"order_date"},
        "Invoices": {"invoice_date", "due_date"},
        "Payments": {"payment_date"},
    }

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        worksheet.row_dimensions[1].height = 24

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        headers = {cell.value: cell.column for cell in worksheet[1]}
        for header in currency_columns.get(worksheet.title, set()):
            for cell in worksheet.iter_cols(
                min_col=headers[header], max_col=headers[header], min_row=2
            ):
                for item in cell:
                    item.number_format = '"₹"#,##0.00'

        for header in date_columns.get(worksheet.title, set()):
            for cell in worksheet.iter_cols(
                min_col=headers[header], max_col=headers[header], min_row=2
            ):
                for item in cell:
                    item.number_format = "yyyy-mm-dd"

        for column_cells in worksheet.columns:
            longest_value = max(len(str(cell.value or "")) for cell in column_cells)
            column_letter = column_cells[0].column_letter
            worksheet.column_dimensions[column_letter].width = min(longest_value + 2, 30)

        for header in currency_columns.get(worksheet.title, set()):
            column_letter = worksheet.cell(row=1, column=headers[header]).column_letter
            worksheet.column_dimensions[column_letter].width = max(
                worksheet.column_dimensions[column_letter].width, 16
            )

        for header in date_columns.get(worksheet.title, set()):
            column_letter = worksheet.cell(row=1, column=headers[header]).column_letter
            worksheet.column_dimensions[column_letter].width = max(
                worksheet.column_dimensions[column_letter].width, 14
            )

    workbook.save(path)


def generate_workbook(output_path: Path = OUTPUT_PATH) -> Path:
    """Write the sample datasets to an Excel workbook and return its path."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    datasets = create_sample_data()

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, dataframe in datasets.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)

    format_workbook(output_path)
    return output_path


if __name__ == "__main__":
    workbook_path = generate_workbook()
    print(f"Created sample workbook: {workbook_path}")
